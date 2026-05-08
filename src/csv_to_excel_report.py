from pathlib import Path
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent.parent
INPUT_FILE = BASE_DIR / "input" / "sample_data.csv"
OUTPUT_DIR = BASE_DIR / "output"


def leggi_transazioni(file_path):
    transazioni = []

    with open(file_path, mode="r", encoding="utf-8") as file:
        reader = csv.DictReader(file)

        for row in reader:
            row["importo"] = float(row["importo"])
            transazioni.append(row)

    return transazioni


def genera_report_excel(transazioni):
    OUTPUT_DIR.mkdir(exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = OUTPUT_DIR / f"report_transazioni_{timestamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Report Transazioni"

    headers = ["Data", "Cliente", "Categoria", "Descrizione", "Importo", "Stato"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for transazione in transazioni:
        ws.append([
            transazione["data"],
            transazione["cliente"],
            transazione["categoria"],
            transazione["descrizione"],
            transazione["importo"],
            transazione["stato"],
        ])

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = '€ #,##0.00'

    ultima_riga = ws.max_row + 2
    totale = sum(t["importo"] for t in transazioni)
    totale_pagato = sum(t["importo"] for t in transazioni if t["stato"] == "Pagato")
    totale_attesa = sum(t["importo"] for t in transazioni if t["stato"] == "In attesa")

    ws[f"D{ultima_riga}"] = "Totale"
    ws[f"E{ultima_riga}"] = totale

    ws[f"D{ultima_riga + 1}"] = "Totale pagato"
    ws[f"E{ultima_riga + 1}"] = totale_pagato

    ws[f"D{ultima_riga + 2}"] = "Totale in attesa"
    ws[f"E{ultima_riga + 2}"] = totale_attesa

    for row in range(ultima_riga, ultima_riga + 3):
        ws[f"D{row}"].font = Font(bold=True)
        ws[f"E{row}"].font = Font(bold=True)
        ws[f"E{row}"].number_format = '€ #,##0.00'

    for column_cells in ws.columns:
        max_length = 0
        column = column_cells[0].column

        for cell in column_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[get_column_letter(column)].width = max_length + 3

    wb.save(output_file)
    return output_file


def main():
    print("Avvio generazione report...")

    transazioni = leggi_transazioni(INPUT_FILE)
    report_path = genera_report_excel(transazioni)

    print(f"Report generato correttamente: {report_path}")


if __name__ == "__main__":
    main()